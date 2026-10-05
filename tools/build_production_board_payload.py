#!/usr/bin/env python3
"""Build a reviewed Dieture Board JSON payload from the production workbook.

The workbook normally includes these two mapping sheets:

* ``User Map``: ``Workbook label`` and ``ERP user email``
* ``Department Map``: ``Workbook department`` and ``ERP Department name``

For a reviewed workbook without those sheets, pass ``--maps-from`` with a
previously reviewed payload.
"""

import argparse
import json
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook


REQUIRED_BET_FIELDS = (
	"short_name", "then_what", "number_today", "number_want", "unit", "by_when", "main_score",
	"so_what", "what", "bet_owner_label", "colour",
)
BET_CONTENT_FIELDS = (
	"short_name", "then_what", "number_today", "number_want", "unit", "so_what", "what",
	"bet_owner_label", "next_step", "tiny_test",
)


def main():
	arguments = _arguments()
	payload = build_payload(
		Path(arguments.workbook),
		version=arguments.version,
		map_payload_path=Path(arguments.maps_from) if arguments.maps_from else None,
		bet_by_when_overrides=_date_overrides(arguments.bet_by_when),
	)
	Path(arguments.output).write_text(json.dumps(payload, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")


def build_payload(workbook_path, *, version=1, map_payload_path=None, bet_by_when_overrides=None):
	workbook = load_workbook(workbook_path, read_only=True, data_only=True)
	user_map, department_map = _maps(workbook, map_payload_path)

	bets = _bets(workbook["Board"], bet_by_when_overrides or {})
	payload = {
		"version": version,
		"user_map": user_map,
		"department_map": department_map,
		"scorecards": _scorecards(workbook["Scorecard"]),
		"bets": bets,
		"promises": _promises(workbook["Promises"], {record["rank"] for record in bets}),
		"standing_numbers": _standing_numbers(workbook["Departments"], workbook["People"]),
		"future_bets": _future_bets(workbook["Future Bets"]),
		"parking_lot": _parking_lot(workbook["Parking Lot"]),
	}
	payload["manifest"] = {section: len(payload[section]) for section in (
		"scorecards", "bets", "promises", "standing_numbers", "future_bets", "parking_lot",
	)}
	return payload


def _scorecards(sheet):
	return [
		{
			"score": _text(row[0]),
			"definition": _text(row[1]),
			"how_measured": _text(row[2]),
			"baseline": _text(row[3]),
			"current": _text(row[4]),
			"target": _text(row[5]),
		}
		for row in _rows(sheet, 6) if _value(row, 0)
	]


def _bets(sheet, by_when_overrides):
	records = []
	for row in _rows(sheet, 10):
		if not isinstance(_value(row, 0), (int, float)):
			continue
		record = {
			"rank": int(_value(row, 0)), "active": 1, "short_name": _text(_value(row, 3)),
			"then_what": _text(_value(row, 4)), "number_today": _text(_value(row, 5)),
			"number_want": _text(_value(row, 6)), "unit": _text(_value(row, 7)), "by_when": _date(_value(row, 8)),
			"main_score": _score(_value(row, 9)), "also_helps": _score(_value(row, 10)),
			"so_what": _text(_value(row, 11)), "so_what_today": _text(_value(row, 12)),
			"so_what_pass_mark": _text(_value(row, 13)), "what": _text(_value(row, 14)),
			"tiny_test": _text(_value(row, 15)), "tiny_test_by": _date(_value(row, 16)),
			"bet_owner_label": _text(_value(row, 17)), "owner_department_label": _text(_value(row, 18)),
			"next_step": _text(_value(row, 19)), "next_step_by": _date(_value(row, 20)),
			"next_step_who_label": _text(_value(row, 21)), "colour": _text(_value(row, 22)),
			"last_updated": _date(_value(row, 23)), "blocker": _text(_value(row, 26)),
			"who_can_unblock": _text(_value(row, 27)), "doors_opened": _text(_value(row, 28)),
		}
		record["by_when"] = by_when_overrides.get(record["rank"], record["by_when"])
		if _is_placeholder_bet(record):
			continue
		_validate_bet(record)
		records.append(record)
	return records


def _promises(sheet, bet_ranks):
	records = []
	for row in _rows(sheet, 6):
		if not _value(row, 0) or int(_value(row, 0)) not in bet_ranks:
			continue
		records.append({
			"bet_rank": int(_value(row, 0)), "deliverable": _text(_value(row, 2)),
			"department_label": _text(_value(row, 3)), "maker_label": _text(_value(row, 4)),
			"due_date": _date(_value(row, 5)), "definition_of_done": _text(_value(row, 8)),
			"receiver_label": _text(_value(row, 9)), "note": _text(_value(row, 20)),
			"workflow": {
				"done": _checkbox(_value(row, 6)), "sent_for_acceptance_on": _date(_value(row, 7)),
				"accepted": _checkbox(_value(row, 11)), "accepted_on": _date(_value(row, 12)),
				"rework_returns": _number(_value(row, 13)), "unclear_returns": _number(_value(row, 14)),
				"last_return_reason": _text(_value(row, 15)),
			},
		})
	return records


def _standing_numbers(department_sheet, people_sheet):
	records = []
	for row in _rows(department_sheet, 6):
		if not _value(row, 0):
			continue
		records.extend(_standing_records(_text(_value(row, 0)), None, None, row, 2, 5))
	for row in _rows(people_sheet, 6):
		if not _value(row, 0):
			continue
		records.extend(_standing_records(
			_text(_value(row, 1)), _text(_value(row, 0)), _text(_value(row, 2)), row, 13, 16,
		))
	return records


def _standing_records(department, person, role_title, row, *starts):
	records = []
	for start in starts:
		metric = _text(_value(row, start))
		if metric and _text(_value(row, start + 1)):
			target = _text(_value(row, start + 1))
			this_week = _text(_value(row, start + 2))
			records.append({
				"department_label": department, "person_label": person, "role_title": role_title,
				"metric": metric, "target": target, "this_week": this_week, "as_of": None,
				"on_target": int(target is not None and target == this_week),
			})
	return records


def _future_bets(sheet):
	return [{
		"title": _text(_value(row, 0)), "unlocked_by_rank": _number(_value(row, 1)),
		"unlocking_bet_name": _text(_value(row, 2)), "unlocking_status": _text(_value(row, 3)),
		"score": _score(_value(row, 4)), "earliest_start": _date(_value(row, 5)),
		"future_owner_label": _text(_value(row, 6)), "why_not_inside": _text(_value(row, 7)),
	} for row in _rows(sheet, 6) if _value(row, 0)]


def _parking_lot(sheet):
	return [{
		"idea": _text(_value(row, 0)), "why_existed": _text(_value(row, 1)),
		"item_owner_label": _text(_value(row, 2)), "parked_on": _date(_value(row, 3)),
		"revisit_by": _date(_value(row, 4)), "still_wanted": _checkbox(_value(row, 6)),
	} for row in _rows(sheet, 6) if _value(row, 0)]


def _maps(workbook, map_payload_path):
	map_sheets = {"User Map", "Department Map"}
	if map_sheets.issubset(workbook.sheetnames):
		return (
			_read_map(workbook, "User Map", "Workbook label", "ERP user email"),
			_read_map(workbook, "Department Map", "Workbook department", "ERP Department name"),
		)
	if not map_payload_path:
		missing = ", ".join(sorted(map_sheets.difference(workbook.sheetnames)))
		raise ValueError("Workbook is missing mapping sheets: {0}".format(missing))
	with map_payload_path.open(encoding="utf-8") as map_file:
		payload = json.load(map_file)
	if not payload.get("user_map") or not payload.get("department_map"):
		raise ValueError("Mapping payload needs user_map and department_map entries.")
	return payload["user_map"], payload["department_map"]


def _read_map(workbook, sheet_name, source_column, target_column):
	if sheet_name not in workbook.sheetnames:
		raise ValueError("Workbook needs a '{0}' sheet.".format(sheet_name))
	rows = list(workbook[sheet_name].iter_rows(values_only=True))
	if not rows:
		raise ValueError("'{0}' is empty.".format(sheet_name))
	headers = {_text(value): index for index, value in enumerate(rows[0]) if _text(value)}
	if source_column not in headers or target_column not in headers:
		raise ValueError("'{0}' needs '{1}' and '{2}' columns.".format(sheet_name, source_column, target_column))
	return {
		_text(row[headers[source_column]]): _text(row[headers[target_column]])
		for row in rows[1:] if _value(row, headers[source_column])
	}


def _rows(sheet, start):
	return sheet.iter_rows(min_row=start, values_only=True)


def _value(row, index):
	return row[index] if len(row) > index else None


def _text(value):
	return str(value).strip() if value not in (None, "") else None


def _date(value):
	if isinstance(value, datetime):
		return value.date().isoformat()
	if isinstance(value, date):
		return value.isoformat()
	if isinstance(value, str):
		try:
			return datetime.fromisoformat(value.strip()).date().isoformat()
		except ValueError:
			return None
	return None


def _number(value):
	return int(value) if isinstance(value, float) and value.is_integer() else value


def _checkbox(value):
	return int(str(value).strip().upper() in {"Y", "YES", "1", "TRUE"})


def _score(value):
	text = _text(value)
	if not text:
		return None
	lower = text.lower()
	if "acqui" in lower:
		return "Acquisition"
	if "retention" in lower:
		return "Retention"
	if "margin" in lower:
		return "Margin"
	return None


def _is_placeholder_bet(record):
	return not any(record.get(field) for field in BET_CONTENT_FIELDS)


def _validate_bet(record):
	missing = [field for field in REQUIRED_BET_FIELDS if not record.get(field)]
	if missing:
		raise ValueError(
			"Board rank {0} is incomplete: {1}".format(record["rank"], ", ".join(missing))
		)


def _date_overrides(values):
	overrides = {}
	for value in values:
		try:
			rank, by_when = value.split("=", 1)
			rank = int(rank)
		except ValueError as exc:
			raise ValueError("--bet-by-when must use RANK=YYYY-MM-DD") from exc
		parsed = _date(by_when)
		if not parsed:
			raise ValueError("--bet-by-when needs an ISO date: {0}".format(value))
		overrides[rank] = parsed
	return overrides


def _arguments():
	parser = argparse.ArgumentParser(description=__doc__)
	parser.add_argument("workbook")
	parser.add_argument("output")
	parser.add_argument("--version", type=int, default=1)
	parser.add_argument("--maps-from", help="JSON payload that provides reviewed User and Department maps")
	parser.add_argument("--bet-by-when", action="append", default=[], metavar="RANK=YYYY-MM-DD")
	return parser.parse_args()


if __name__ == "__main__":
	main()
